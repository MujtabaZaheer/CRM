import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_page_number_to_run(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_border_to_image_table(cell):
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def build_report():
    doc = Document()

    # Base Page setup (1 inch margins, standard Letter size)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.different_first_page_header_footer = True

        # Footer for subsequent pages: Right aligned page number
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = p_foot.add_run()
        f_run.font.name = 'Times New Roman'
        f_run.font.size = Pt(10)
        f_run.font.color.rgb = RGBColor(100, 100, 100)
        add_page_number_to_run(f_run)

    # Base style definitions
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11.5)
    normal_style.font.color.rgb = RGBColor(20, 20, 20)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # ----------------------------------------------------
    # 1. TITLE / COVER PAGE
    # ----------------------------------------------------
    p_top_spacer = doc.add_paragraph()
    p_top_spacer.paragraph_format.space_before = Pt(28)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(32)
    r_title = p_title.add_run("Summer Internship Report")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # CUST Logo
    logo_path = os.path.join(os.getcwd(), 'extracted_images', 'page_1_img_1_Image13.png')
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(36)
        r_logo = p_logo.add_run()
        r_logo.add_picture(logo_path, width=Inches(2.0))
    else:
        p_logo = doc.add_paragraph()
        p_logo.paragraph_format.space_after = Pt(36)

    # Submitted by
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(3)
    r_sub = p_sub.add_run("Submitted by :")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(12.5)

    p_student = doc.add_paragraph()
    p_student.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_student.paragraph_format.space_after = Pt(26)
    r_student = p_student.add_run("SAAD ALI QURESHI-------------BSE223018")
    r_student.font.name = 'Times New Roman'
    r_student.font.size = Pt(13.5)
    r_student.font.bold = True

    # Submitted to
    p_to = doc.add_paragraph()
    p_to.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_to.paragraph_format.space_after = Pt(3)
    r_to = p_to.add_run("Submitted to :")
    r_to.font.name = 'Times New Roman'
    r_to.font.size = Pt(12.5)

    p_sir = doc.add_paragraph()
    p_sir.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sir.paragraph_format.space_after = Pt(26)
    r_sir = p_sir.add_run("Dr. Nadeem Anjum (HOD)")
    r_sir.font.name = 'Times New Roman'
    r_sir.font.size = Pt(14)
    r_sir.font.bold = True

    # Duration
    p_dur = doc.add_paragraph()
    p_dur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dur.paragraph_format.space_after = Pt(32)
    r_dur = p_dur.add_run("Internship Duration: 3 August, 2026 – 22 September, 2026 (7 Weeks)")
    r_dur.font.name = 'Times New Roman'
    r_dur.font.size = Pt(12.5)
    r_dur.font.bold = True

    # Department & University
    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_after = Pt(4)
    r_dept = p_dept.add_run("Department Of Software Engineering")
    r_dept.font.name = 'Times New Roman'
    r_dept.font.size = Pt(13)

    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_univ = p_univ.add_run("Capital University Of Science & Technology, Islamabad")
    r_univ.font.name = 'Times New Roman'
    r_univ.font.size = Pt(13)

    doc.add_page_break()

    # ----------------------------------------------------
    # 2. TABLE OF CONTENTS
    # ----------------------------------------------------
    p_toc_heading = doc.add_paragraph()
    p_toc_heading.paragraph_format.space_before = Pt(12)
    p_toc_heading.paragraph_format.space_after = Pt(24)
    r_th = p_toc_heading.add_run("Table of Contents")
    r_th.font.name = 'Times New Roman'
    r_th.font.size = Pt(16)
    r_th.font.bold = True

    toc_entries = [
        ("Acknowledgement", "3", False),
        ("Introduction", "4", False),
        ("DAILY ASSIGNMENT TASKS", "5", True),
        ("Week 1: Foundations, Architecture & Multi-Role Authentication", "5", False),
        ("Week 2: Student Experience, 4-Stage Onboarding & Catalogs", "8", False),
        ("Week 3: Counsellor Workspace, Lead Ingestion & Routing Engine", "12", False),
        ("Week 4: Application Pipeline, Admissions & Document Vault", "16", False),
        ("Week 5: Agent Partner Portal, University Desk & Finance Engine", "20", False),
        ("Week 6: AI Intelligence Suite, Gemini OCR & Quality Monitors", "24", False),
        ("Week 7: Automated Testing, Optimization & Final Handover", "28", False),
        ("Conclusion", "32", True),
        ("References", "33", True),
    ]

    for title, page_str, is_bold in toc_entries:
        p_row = doc.add_paragraph()
        p_row.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p_row.paragraph_format.space_after = Pt(5)
        p_row.paragraph_format.line_spacing = 1.15
        
        r1 = p_row.add_run(title)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        r1.font.bold = is_bold
        
        r2 = p_row.add_run(f"\t{page_str}")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        r2.font.bold = is_bold

    doc.add_page_break()

    # ----------------------------------------------------
    # 3. ACKNOWLEDGEMENT
    # ----------------------------------------------------
    p_ack_head = doc.add_paragraph()
    p_ack_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ack_head.paragraph_format.space_before = Pt(18)
    p_ack_head.paragraph_format.space_after = Pt(28)
    r_ah = p_ack_head.add_run("Acknowledgement")
    r_ah.font.name = 'Times New Roman'
    r_ah.font.size = Pt(16)
    r_ah.font.bold = True

    p_ack0 = doc.add_paragraph()
    p_ack0.paragraph_format.line_spacing = 1.25
    p_ack0.paragraph_format.space_after = Pt(14)
    p_ack0.add_run(
        "First and foremost, I would like to express my deepest gratitude and profound respect to my academic supervisor, "
        "Dr. Nadeem Anjum (HOD, Department of Software Engineering, Capital University of Science & Technology). Under his continuous "
        "academic supervision, visionary direction, and rigorous engineering guidance, this summer internship was successfully planned and executed. "
        "On the first day of my internship, Dr. Nadeem Anjum established the comprehensive project scope, task breakdown, architectural standards, "
        "and daily milestone review framework. Every day following the completion of my technical work, I submitted my completed modules, code artifacts, "
        "and dashboard deliverables to Dr. Nadeem Anjum for thorough inspection and validation. His timely feedback, high engineering standards, and "
        "continuous verification of each milestone before authorizing progression to the next task provided the structural backbone for my learning."
    )

    p_ack1 = doc.add_paragraph()
    p_ack1.paragraph_format.line_spacing = 1.25
    p_ack1.paragraph_format.space_after = Pt(14)
    p_ack1.add_run(
        "I would also like to express my sincere gratitude to the Director of Convo (Pvt.) Ltd., Mr. Sanaullah Irfan, for providing me "
        "with the opportunity to complete my full-stack software engineering internship at the organization's head office in Islamabad. "
        "I am deeply thankful for his visionary architectural guidance and the high-performance professional environment fostered at Convo. "
        "His insights into multi-tenant software systems, enterprise cloud design, and agentic AI integrations greatly enriched my industrial experience."
    )

    p_ack2 = doc.add_paragraph()
    p_ack2.paragraph_format.line_spacing = 1.25
    p_ack2.paragraph_format.space_after = Pt(14)
    p_ack2.add_run(
        "I also extend my heartfelt appreciation to my team lead at Convo, Mr. Waqar Akram, for his relentless technical mentorship, "
        "constructive code reviews, and practical guidance throughout the design, development, and testing of the Education CRM (EduCRM) platform. "
        "His recommendations on state management, asynchronous data flows, and automated testing were instrumental to my day-to-day progress."
    )

    p_ack3 = doc.add_paragraph()
    p_ack3.paragraph_format.line_spacing = 1.25
    p_ack3.paragraph_format.space_after = Pt(45)
    p_ack3.add_run(
        "This internship has played an essential role in bridging theoretical concepts taught at Capital University of Science & Technology (CUST) "
        "with real-world software engineering practices, instilling in me the confidence, problem-solving discipline, and technical depth required to "
        "excel as a software engineering professional."
    )

    p_sig_hod = doc.add_paragraph()
    p_sig_hod.paragraph_format.space_after = Pt(3)
    r_sh = p_sig_hod.add_run("DR. NADEEM ANJUM (HOD)")
    r_sh.font.bold = True
    r_sh.font.size = Pt(12)

    p_sig_hod_sub = doc.add_paragraph()
    p_sig_hod_sub.paragraph_format.space_after = Pt(16)
    r_shs = p_sig_hod_sub.add_run("Head, Department of Software Engineering, CUST")
    r_shs.font.size = Pt(11)

    p_sig_hod_line = doc.add_paragraph()
    p_sig_hod_line.paragraph_format.space_after = Pt(28)
    r_shl = p_sig_hod_line.add_run("HOD Signature & Stamp : _________________________________")
    r_shl.font.size = Pt(11)

    p_sig1 = doc.add_paragraph()
    p_sig1.paragraph_format.space_after = Pt(3)
    r_s1 = p_sig1.add_run("SANAULLAH IRFAN")
    r_s1.font.bold = True
    r_s1.font.size = Pt(12)

    p_sig2 = doc.add_paragraph()
    p_sig2.paragraph_format.space_after = Pt(16)
    r_s2 = p_sig2.add_run("(Director Engineering & AI & Agentic Technology, Convo Pvt. Ltd.)")
    r_s2.font.size = Pt(11)

    p_sig3 = doc.add_paragraph()
    r_s3 = p_sig3.add_run("Industry Mentor Signature : _________________________________")
    r_s3.font.size = Pt(11)

    doc.add_page_break()

    # ----------------------------------------------------
    # 4. INTRODUCTION
    # ----------------------------------------------------
    p_intro_h = doc.add_paragraph()
    p_intro_h.paragraph_format.space_before = Pt(12)
    p_intro_h.paragraph_format.space_after = Pt(16)
    r_ih = p_intro_h.add_run("Introduction")
    r_ih.font.name = 'Times New Roman'
    r_ih.font.size = Pt(16)
    r_ih.font.bold = True

    p_intro_body1 = doc.add_paragraph()
    p_intro_body1.paragraph_format.line_spacing = 1.2
    p_intro_body1.paragraph_format.space_after = Pt(12)
    p_intro_body1.add_run(
        "Convo (Pvt.) Ltd. is an innovative software consulting and technology firm specializing in full-stack development, "
        "cloud-based SaaS solutions, enterprise collaboration platforms, and agentic artificial intelligence systems. "
        "Focused on engineering robust, high-performance web and mobile platforms, the organization leverages cutting-edge product "
        "development lifecycles and modern software architectures to streamline complex business workflows, enhance operational transparency, "
        "and accelerate digital transformation for global clients."
    )

    # Convo banner
    banner_path = os.path.join(os.getcwd(), 'extracted_images', 'page_19_img_1_Image324.png')
    if os.path.exists(banner_path):
        p_banner = doc.add_paragraph()
        p_banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_banner.paragraph_format.space_after = Pt(16)
        r_ban = p_banner.add_run()
        r_ban.add_picture(banner_path, width=Inches(6.0))

    p_intro_body2 = doc.add_paragraph()
    p_intro_body2.paragraph_format.line_spacing = 1.2
    p_intro_body2.paragraph_format.space_after = Pt(12)
    p_intro_body2.add_run(
        "During my 7-week summer internship (3 August, 2026 to 22 September, 2026), conducted under the academic supervision of "
        "Dr. Nadeem Anjum (HOD, Department of Software Engineering, CUST) and industry mentorship at Convo (Pvt.) Ltd., I was assigned "
        "to architect and develop core modules for the Education CRM and Student Application Management Platform (EduCRM). "
        "EduCRM is an enterprise-grade, multi-tenant web ecosystem tailored for international education consultants, global university "
        "recruitment networks, counsellors, and prospective students. The platform replaces fragmented, spreadsheet-driven operations "
        "with an integrated, intelligent, and automated operational backbone."
    )

    p_intro_body3 = doc.add_paragraph()
    p_intro_body3.paragraph_format.line_spacing = 1.2
    p_intro_body3.paragraph_format.space_after = Pt(12)
    p_intro_body3.add_run(
        "Key capabilities of the EduCRM ecosystem include prospective-student lead ingestion and deduplication, dynamic multi-branch lead routing, "
        "a 4-stage interactive student onboarding experience, international university and programme discovery catalogs, automated "
        "document vault compliance with zero-cost cloud storage, university application pipeline tracking, external agent referral portals "
        "with automated commission ledgers, and deep Google Gemini AI integrations for automated CV/resume parsing and intelligent programme eligibility matching."
    )

    p_tech = doc.add_paragraph()
    p_tech.paragraph_format.line_spacing = 1.2
    p_tech.paragraph_format.space_after = Pt(20)
    p_tech.add_run(
        "The application is built on a modern high-performance technology stack comprising React 19, TypeScript, Vite, Tailwind CSS v4, "
        "Firebase Firestore for real-time document-oriented data persistence, and Playwright for comprehensive end-to-end automated testing."
    )

    # ----------------------------------------------------
    # 5. DAILY ASSIGNMENT TASKS - WEEKS 1 TO 7
    # ----------------------------------------------------
    p_tasks_h = doc.add_paragraph()
    p_tasks_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tasks_h.paragraph_format.space_before = Pt(16)
    p_tasks_h.paragraph_format.space_after = Pt(18)
    r_th = p_tasks_h.add_run("DAILY ASSIGNMENT TASKS")
    r_th.font.name = 'Times New Roman'
    r_th.font.size = Pt(16)
    r_th.font.bold = True

    # Helper function for adding day entry
    def add_day_entry(date_str, content):
        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_before = Pt(10)
        p_d.paragraph_format.space_after = Pt(4)
        r_d = p_d.add_run(date_str)
        r_d.font.bold = True
        r_d.font.size = Pt(11.5)

        p_c = doc.add_paragraph()
        p_c.paragraph_format.line_spacing = 1.2
        p_c.paragraph_format.space_after = Pt(10)
        r_c = p_c.add_run(content)
        r_c.font.size = Pt(11)

    # Helper function for inserting image with border and caption (scaled to 6.0 inches)
    def add_image_box(img_path, caption_text, width_inch=6.0):
        if not os.path.exists(img_path):
            print(f"Warning: Image not found: {img_path}")
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        add_border_to_image_table(cell)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run()
        r.add_picture(img_path, width=Inches(width_inch))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(14)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(90, 90, 90)

    # Helper function for week header
    def add_week_header(week_title):
        p_w = doc.add_paragraph()
        p_w.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_w.paragraph_format.space_before = Pt(20)
        p_w.paragraph_format.space_after = Pt(14)
        r_w = p_w.add_run(week_title)
        r_w.font.name = 'Times New Roman'
        r_w.font.size = Pt(14.5)
        r_w.font.bold = True
        r_w.font.color.rgb = RGBColor(15, 23, 42)

    # ----------------------------------------------------
    # WEEK 1: Foundations, Architecture & Multi-Role Authentication
    # ----------------------------------------------------
    add_week_header("Week 1: Foundations, Architecture & Multi-Role Authentication")

    add_day_entry(
        "3 August, 2026 — Monday",
        "On the first day of my internship, I arrived at Capital University of Science & Technology (CUST) to meet with the Head of the "
        "Department of Software Engineering, Dr. Nadeem Anjum (HOD), under whose supervision my internship was conducted. "
        "Dr. Nadeem Anjum provided the complete internship roadmap, milestone structure, technical objectives, and daily reporting protocol. "
        "He established a strict review workflow: each day after completing my software engineering tasks, I would submit the deliverables, "
        "code implementations, and UI screens to Dr. Nadeem Anjum for detailed review and evaluation. Upon his assessment and approval, I would proceed "
        "to the subsequent task. Following this briefing, I arrived at the head office of Convo (Pvt.) Ltd. in Islamabad for my full-stack software "
        "engineering internship. I attended a comprehensive orientation session conducted by HR and engineering management detailing the organization's "
        "culture, product portfolio, and engineering standards. During the session, I was introduced to the project I would be contributing to: "
        "the multi-tenant Education CRM and Student Application Management Platform (EduCRM). I set up my local development environment by configuring "
        "Node.js v22, Git version control, Visual Studio Code extensions, and cloning the project repository. I verified initial dependencies including "
        "React 19, Vite 6, Tailwind CSS v4, and Firebase client SDKs. At the end of the day, I prepared my first daily task report and submitted it to "
        "Dr. Nadeem Anjum, who evaluated the setup and authorized commencement of architecture exploration."
    )

    add_day_entry(
        "4 August, 2026 — Tuesday",
        "On my second day, I had an architectural strategy session with the Director of Convo, Mr. Sanaullah Irfan. "
        "He explained the overarching vision behind EduCRM, highlighting how traditional education consulting firms struggle with fragmented "
        "communication, manual paper-based application pipelines, and high customer acquisition costs. He detailed the architectural "
        "necessity of multi-tenancy, strict data isolation per organization, and granular role-based access control (RBAC). Following "
        "his session, I explored the project directory structure, examining how modular routes and contexts are structured to support diverse "
        "stakeholder roles including Super Admins, Office Managers, Counsellors, Admissions Officers, External Agents, and Students. "
        "I documented the module taxonomy and submitted the findings to Dr. Nadeem Anjum (HOD), who verified the design integrity."
    )

    add_day_entry(
        "5 August, 2026 — Wednesday",
        "On my third day, I met with my team lead, Mr. Waqar Akram, who provided an in-depth walkthrough of the technology stack and "
        "state management architecture. He explained how Firestore collections are organized hierarchically with tenant and office scoping, "
        "and how reactive listeners synchronize live updates across counselling desks. To build hands-on familiarity, he assigned me practical "
        "tasks: writing modular React components using TypeScript interfaces, configuring route guards, and practicing mock session injection "
        "for automated browser testing. I also reviewed Firestore security rule structures in firestore.rules to understand read/write restrictions "
        "enforced on user roles. I submitted the security rules audit to Dr. Nadeem Anjum (HOD) for technical sign-off."
    )

    add_day_entry(
        "6 August, 2026 — Thursday",
        "On the fourth day, I focused on implementing the unified Authentication and Account Access module (/login). Under Mr. Waqar Akram's "
        "supervision, I designed and coded a clean, responsive authentication interface equipped with email/password validation, interactive "
        "error alerts for invalid credentials, and a Fast Role Demo Switcher. This switcher enables developers, testers, and stakeholders to "
        "instantly simulate logins as a Student, Counsellor, Admissions Officer, External Agent, or Super Admin without repeatedly typing credentials. "
        "I ensured that HTML5 client-side validation rules handled empty inputs and formatted error messages gracefully. The working authentication "
        "portal was submitted to Dr. Nadeem Anjum (HOD), who inspected the demo switching capability and approved the deliverable."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '01_login_page.png'),
        "Figure 1.1: EduCRM Unified Authentication Portal with Fast Role Demo Switcher"
    )

    add_day_entry(
        "7 August, 2026 — Friday",
        "On the fifth day, I implemented robust session persistence and route protection mechanisms. I developed local storage state "
        "hydration logic within AuthContext.tsx, ensuring that user sessions persist across browser page refreshes and window reloads. "
        "I also configured the email verification gate (/verify-email) and protected route guards (ProtectedLayout.tsx, RoleRoute.tsx), "
        "preventing unauthenticated users or unauthorized roles from accessing restricted platform areas. I rendered and verified the formal "
        "sequence diagram for the multi-tenant authentication and session impersonation flow. At the end of the week, I demonstrated "
        "the authentication workflows and role-switching capabilities to both Mr. Waqar Akram and Dr. Nadeem Anjum (HOD), receiving full approval."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'diagrams_png', 'SD1_Authentication_and_Impersonation.png'),
        "Figure 1.2: Sequence Diagram — Multi-Tenant Authentication & Session Impersonation Flow"
    )

    # ----------------------------------------------------
    # WEEK 2: Student Experience, 4-Stage Onboarding & Catalogs
    # ----------------------------------------------------
    doc.add_page_break()
    add_week_header("Week 2: Student Experience, 4-Stage Onboarding & Catalogs")

    add_day_entry(
        "10 August, 2026 — Monday",
        "At the start of the second week, I attended an architectural planning session with Director Mr. Sanaullah Irfan focused on "
        "student user journeys and user-centric frontend design. He stressed that prospective students require a frictionless, visually engaging, "
        "and reassuring experience when navigating international university applications. Guided by this directive, I planned the architecture "
        "for the prospective student experience portal, breaking down the application intake into an intuitive 4-stage progressive onboarding "
        "wizard: personal profile, academic background, study preferences, and document submission checklist. I submitted the onboarding wireframes "
        "and step state machine to Dr. Nadeem Anjum (HOD) for academic evaluation."
    )

    add_day_entry(
        "11 August, 2026 — Tuesday",
        "Today, I began developing Onboarding Stage 1 (/student/onboarding/stage-1) and Stage 2 (/student/onboarding/stage-2). "
        "Stage 1 captures essential identity attributes including full legal name, date of birth, passport details, nationality, and emergency "
        "contact information. In Stage 2, I engineered dynamic form arrays allowing applicants to input multiple academic qualifications "
        "(high school, undergraduate degrees, grading systems, and GPAs) alongside standardized language proficiency scores (IELTS, TOEFL, PTE, Duolingo). "
        "I implemented real-time validation schemas to prevent malformed score submissions. The validated onboarding stages were presented to "
        "Dr. Nadeem Anjum (HOD) for technical review."
    )

    add_day_entry(
        "12 August, 2026 — Wednesday",
        "I continued developing the onboarding wizard by building Stage 3 (Study Preferences) and Stage 4 (Document Checklist Vault). "
        "Stage 3 captures the student's target destinations (United Kingdom, United States, Canada, Australia, Ireland), desired study levels "
        "(Bachelors, Masters, PhD), target intake periods (Fall 2026, Spring 2027), and annual budget thresholds. In Stage 4, I constructed a "
        "responsive document upload checklist notifying students of required credentials—such as academic transcripts, passport scans, letters of "
        "recommendation, and Statements of Purpose (SOP)—and persisting onboarding completion flags in Firestore. I submitted the complete 4-stage "
        "flow to Dr. Nadeem Anjum (HOD) who commended the comprehensive data model."
    )

    add_day_entry(
        "13 August, 2026 — Thursday",
        "On the fourth day of Week 2, I developed the primary Student Experience Dashboard (/student/dashboard). The dashboard serves as "
        "the central command center for applicants, featuring a visual Application Progress Timeline that displays active pipeline stages "
        "(Submitted, Document Verification, Offer Issued, Visa Processing). I integrated actionable metric cards showing application counts, "
        "pending documentation alerts, and an assigned counsellor contact card enabling direct messaging. I styled the interface using Tailwind CSS, "
        "incorporating smooth transition animations and responsive cards. I submitted the dashboard implementation to Dr. Nadeem Anjum (HOD) for inspection."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '05_student_dashboard.png'),
        "Figure 2.1: Student Experience Dashboard — Application Timeline & Metric Cards"
    )

    add_day_entry(
        "14 August, 2026 — Friday",
        "To conclude the week, I engineered the Global University Discovery Catalog (/student/universities) and Programme Search Engine "
        "(/student/programs). I built multi-attribute filter components allowing prospective students to explore partner institutions by "
        "country, world ranking, tuition fee range, campus locations, and language requirements. Each university card features campus imagery, "
        "accreditations, and an interactive 'Explore Programmes' link. I tested the catalog responsiveness across mobile and desktop breakpoints "
        "and submitted the complete module to Dr. Nadeem Anjum (HOD), who verified the query filters and approved the week's milestones."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '06_university_catalog.png'),
        "Figure 2.2: Global University Discovery Catalog with Destination & Attribute Filters"
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '07_programme_search.png'),
        "Figure 2.3: Academic Programme Discovery Engine with Fee & Intake Filters"
    )

    # ----------------------------------------------------
    # WEEK 3: Counsellor Workspace, Lead Ingestion & Routing Engine
    # ----------------------------------------------------
    doc.add_page_break()
    add_week_header("Week 3: Counsellor Workspace, Lead Ingestion & Routing Engine")

    add_day_entry(
        "17 August, 2026 — Monday",
        "At the beginning of the third week, Director Mr. Sanaullah Irfan led an interactive session on lead lifecycle management, "
        "counselling operations, and sales pipeline optimization in educational recruitment. He explained how leads progress through strict "
        "lifecycle stages (New, Contacted, Qualified, Application In-Progress, Offer Received, Visa Lodged, Enrolled). He highlighted the danger "
        "of lead leakage and emphasized the need for automated lead distribution rules, deduplication mechanisms, and real-time activity logging. "
        "Following the session, I mapped out the counsellor workspace architecture and submitted the specifications to Dr. Nadeem Anjum (HOD)."
    )

    add_day_entry(
        "18 August, 2026 — Tuesday",
        "Today, I developed the Counsellor Pipeline Dashboard (/counsellor/dashboard). This interface empowers academic counsellors "
        "to monitor their assigned student portfolio, track urgent follow-ups, and review daily operational metrics. I integrated summary widgets "
        "displaying Total Assigned Leads, Active Applications, Documents Requiring Verification, and Pending Visa Submissions. I also added a "
        "priority action list highlighting students with approaching university application deadlines, ensuring timely follow-up communication. "
        "I submitted the completed counsellor dashboard to Dr. Nadeem Anjum (HOD) for technical validation."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '02_counsellor_dashboard.png'),
        "Figure 3.1: Counsellor Pipeline Dashboard — Workload Metrics & Priority Follow-ups"
    )

    add_day_entry(
        "19 August, 2026 — Wednesday",
        "I focused on engineering the Global Leads Management Desk (/leads). I implemented a high-performance data table displaying "
        "lead records across all operational branches. I added multi-column sorting, search filters by applicant name, email, country preference, "
        "and acquisition channel (organic web signup, social campaign, external agent referral, walk-in inquiry). I also created a modal dialog "
        "allowing counsellors to manually ingest new walk-in leads with instant validation and branch assignment. The lead management console "
        "was submitted to Dr. Nadeem Anjum (HOD) for functional review."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '03_leads_management.png'),
        "Figure 3.2: Global Leads Management Desk — Search Filters & Ingestion Console"
    )

    add_day_entry(
        "20 August, 2026 — Thursday",
        "Today, I implemented the automated Lead Routing Engine (/lead-routing) and lead deduplication logic. Working with Mr. Waqar Akram, "
        "I constructed rule-based distribution algorithms that automatically assign incoming student leads to available branch counsellors using "
        "a configurable round-robin strategy and geographic specialization (e.g., routing UK applicants to UK specialist counsellors). I also built "
        "deduplication filters that detect existing email addresses and phone numbers upon lead entry, preventing duplicate records from diluting "
        "counsellor workloads. I rendered the formal sequence diagram and submitted the engine to Dr. Nadeem Anjum (HOD) for review."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '13_lead_routing_rules.png'),
        "Figure 3.3: Automated Lead Routing & Distribution Configuration Console"
    )

    add_image_box(
        os.path.join(os.getcwd(), 'diagrams_png', 'SD2_Lead_Lifecycle_and_Deduplication.png'),
        "Figure 3.4: Sequence Diagram — Automated Lead Lifecycle, Deduplication & Routing Flow"
    )

    add_day_entry(
        "21 August, 2026 — Friday",
        "On the final day of Week 3, I developed the Assigned Students Directory (/counsellor/students) and integrated an internal "
        "messaging console enabling direct real-time communication between counsellors and students. I implemented student profile inspection "
        "drawers displaying academic records, uploaded transcripts, and past counselling notes. In the afternoon, I held a weekly progress review "
        "with Mr. Waqar Akram and submitted the deliverables to Dr. Nadeem Anjum (HOD). Dr. Nadeem Anjum commended the modularity of the routing "
        "rules engine and approved the completion of Week 3."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '18_assigned_students.png'),
        "Figure 3.5: Assigned Students Directory & Active Case Dossier Console"
    )

    # ----------------------------------------------------
    # WEEK 4: Application Pipeline, Admissions & Document Vault
    # ----------------------------------------------------
    doc.add_page_break()
    add_week_header("Week 4: Application Pipeline, Admissions & Document Vault")

    add_day_entry(
        "24 August, 2026 — Monday",
        "At the start of the fourth week, I conducted a technical review of international university admissions workflows with my team lead, "
        "Mr. Waqar Akram. We analyzed university requirements across key study destinations including the United Kingdom (UCAS and direct partner "
        "portals), the United States (Common App and institutional systems), and Australia (PRISMS and direct agent interfaces). We established "
        "the standardized state machine for applications within EduCRM: Draft, Under Review, Submitted to University, Conditional Offer, "
        "Unconditional Offer, CAS / COE Issued, Visa Processing, and Final Enrolment. I submitted the state transition schema to Dr. Nadeem Anjum (HOD)."
    )

    add_day_entry(
        "25 August, 2026 — Tuesday",
        "Today, I developed the core University Applications Tracking Pipeline (/applications). I engineered a unified kanban and table view "
        "allowing admissions officers and counsellors to track student applications across partner universities. I built status badge components "
        "with distinct color codes for offer types and deadline urgency. I also implemented application creation workflows allowing counsellors "
        "to link verified student profiles directly to target university programmes with pre-populated academic transcripts and recommendation letters. "
        "The pipeline tracking console was submitted to Dr. Nadeem Anjum (HOD) for verification."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '04_applications_pipeline.png'),
        "Figure 4.1: Multi-Stage University Applications Tracking Pipeline Desk"
    )

    add_day_entry(
        "26 August, 2026 — Wednesday",
        "I developed the Admissions Desk Verification console (/admissions/verification). This interface is dedicated to admissions "
        "officers responsible for vetting applicant credentials before institutional submission. I implemented interactive document preview "
        "tools, verification checkboxes for academic certificates and English language proofs, and conditional offer logging. When an applicant "
        "receives an institutional offer, the officer can record offer conditions, deposit deadlines, and fee concessions directly into the student dossier. "
        "I submitted the verification console to Dr. Nadeem Anjum (HOD) who validated the role gating logic."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '14_admissions_verification.png'),
        "Figure 4.2: Admissions Desk Credential Verification & Offer Console"
    )

    add_day_entry(
        "27 August, 2026 — Thursday",
        "Today, I focused on implementing the Zero-Cost Cloud Document Storage architecture. Due to high cloud storage costs associated with "
        "large student dossiers (scanned transcripts, portfolios, financial affidavits, bank statements), Convo devised a hybrid Google Drive storage "
        "model. I implemented client-side file upload handlers that stream binary payloads to Google Drive storage via Apps Script endpoints, "
        "storing secure, signed file IDs and URLs in Firestore. I added robust fallback logic to store metadata securely even in offline or low-bandwidth scenarios. "
        "I rendered the storage sequence diagram and presented the architecture to Dr. Nadeem Anjum (HOD) for technical evaluation."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'diagrams_png', 'SD4_Zero_Cost_Google_Drive_Upload.png'),
        "Figure 4.3: Architecture Diagram — Zero-Cost Cloud Storage & Document Ingestion Flow"
    )

    add_day_entry(
        "28 August, 2026 — Friday",
        "To wrap up the fourth week, I finalized the Student Document Vault (/student/documents) and tested the end-to-end application lifecycle. "
        "I ensured that when an admissions officer marks a document as 'Verified' or 'Requires Re-upload', automated status events are dispatched "
        "to the student portal with real-time notifications. I also validated immutable audit logging for all application state changes. I rendered "
        "the application processing sequence diagram. I submitted the week's deliverables to Dr. Nadeem Anjum (HOD), who verified the document "
        "lifecycle and approved advancement to Week 5."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '08_student_documents.png'),
        "Figure 4.4: Student Document Vault & Compliance Console"
    )

    add_image_box(
        os.path.join(os.getcwd(), 'diagrams_png', 'SD5_Application_Pipeline_and_Admissions.png'),
        "Figure 4.5: Sequence Diagram — University Application Processing & Admissions Desk Flow"
    )

    # ----------------------------------------------------
    # WEEK 5: Agent Partner Portal, University Desk & Finance Engine
    # ----------------------------------------------------
    doc.add_page_break()
    add_week_header("Week 5: Agent Partner Portal, University Desk & Finance Engine")

    add_day_entry(
        "31 August, 2026 — Monday",
        "At the start of the fifth week, Director Mr. Sanaullah Irfan held an insightful session on B2B partner ecosystems, recruitment "
        "sub-agent networks, and financial reconciliation in international education consulting. He explained that a modern CRM must not only "
        "serve internal staff but also empower external stakeholders—such as localized recruitment agents and university partner admissions representatives—"
        "while maintaining strict multi-tenant isolation, data privacy, and role boundaries. I outlined technical requirements for the Agent Partner "
        "Portal and the University Partner Portal and submitted them to Dr. Nadeem Anjum (HOD) for academic review."
    )

    add_day_entry(
        "1 September, 2026 — Tuesday",
        "Today, I built the External Agent Partner Portal (/agent/dashboard). External education agents refer prospective students to consulting "
        "firms and require full visibility into the admission progress of their referrals. I engineered the Agent Referral Submission form "
        "(/agent/refer-lead), which allows agents to register prospective students with attached qualifications. I also built the Referrals Ledger "
        "(/agent/referrals), giving agents real-time transparency into whether their referred students have received university offers, paid deposits, "
        "or enrolled. I submitted the portal to Dr. Nadeem Anjum (HOD) for inspection."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '09_agent_portal.png'),
        "Figure 5.1: External Agent Partner Portal — Student Referrals & Commission Ledger"
    )

    add_day_entry(
        "2 September, 2026 — Wednesday",
        "I developed the Agent Commission Ledger and Statement Generator (/agent/commissions). In education recruitment, agents earn "
        "contractual commissions based on student university enrolments. I programmed automated calculation algorithms supporting percentage-based "
        "and flat-rate commission tiers, milestone release conditions (e.g., 50% upon initial tuition deposit, 50% upon post-arrival census date), "
        "and downloadable billing statements. This eliminated manual spreadsheet reconciliations between finance officers and external agent networks. "
        "I rendered the agent referral sequence diagram and submitted it to Dr. Nadeem Anjum (HOD) for technical verification."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'diagrams_png', 'SD6_Agent_Network_and_Commissions.png'),
        "Figure 5.2: Sequence Diagram — External Agent Referral Lifecycle & Commission Calculation"
    )

    add_day_entry(
        "3 September, 2026 — Thursday",
        "Today, I engineered the University Partner Portal (/university/dashboard and /university/applications). This portal provides "
        "designated admissions officers from partner universities with secure, read-and-update access to applications submitted specifically to their "
        "institution. I implemented an application review console where university representatives can directly issue conditional offers, upload "
        "official acceptance letters, and manage Confirmation of Acceptance for Studies (CAS) / Confirmation of Enrolment (COE) issuance records "
        "(/university/cas-issuance). The implementation was submitted to Dr. Nadeem Anjum (HOD) who commended the multi-stakeholder design."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '10_university_partner_portal.png'),
        "Figure 5.3: University Partner Portal — Direct Application Review & CAS Issuance Console"
    )

    add_day_entry(
        "4 September, 2026 — Friday",
        "To conclude the fifth week, I worked on the Finance and Invoicing Engine (/finance/invoices, /finance/payments). I built "
        "interfaces for tracking student service fees, university commission receivables, and student refund requests. I also integrated "
        "the Visa Officer workspace (/visa-officer/dashboard) to track visa appointment dates, document checklists (financial statements, TB tests, "
        "biometrics), and visa grant/refusal outcomes. I conducted an extensive review of the B2B partner and finance features with Mr. Waqar Akram "
        "and submitted the completed weekly dossier to Dr. Nadeem Anjum (HOD), who authorized proceeding to AI intelligence integrations."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '17_finance_invoices.png'),
        "Figure 5.4: Finance & Invoicing Engine — Student Invoicing, Payments & Commission Ledger"
    )

    # ----------------------------------------------------
    # WEEK 6: AI Intelligence Suite, Gemini OCR & Quality Monitors
    # ----------------------------------------------------
    doc.add_page_break()
    add_week_header("Week 6: AI Intelligence Suite, Gemini OCR & Quality Monitors")

    add_day_entry(
        "7 September, 2026 — Monday",
        "At the start of the sixth week, Director Mr. Sanaullah Irfan led an advanced session on integrating agentic AI capabilities and large "
        "language models into enterprise business software. He emphasized that in the modern AI era, software engineering is shifting from manual "
        "data entry toward AI-augmented intelligence. Under his direction, I worked on integrating Google Gemini AI into the EduCRM platform. "
        "I implemented the automated Student CV/Resume OCR Extraction component (StudentCVUploader), which parses uploaded student resumes and "
        "automatically extracts education histories, GPA metrics, language test scores, and work experience directly into structured Firestore documents. "
        "I rendered the Gemini AI sequence diagram and submitted the architecture to Dr. Nadeem Anjum (HOD)."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'diagrams_png', 'SD3_Gemini_AI_Counsellor_Suite.png'),
        "Figure 6.1: Sequence Diagram — Google Gemini AI Resume OCR & Counsellor Suite Integration"
    )

    add_day_entry(
        "8 September, 2026 — Tuesday",
        "Today, I developed the AI Intelligent Programme Eligibility Matcher (eligibility.ts) and the Lead Scoring Matrix configuration "
        "engine (/lead-scoring). The eligibility engine algorithmically compares a student's parsed qualifications and English scores against "
        "institutional entry criteria, computing real-time compatibility scores and flagging admission gaps. I also built the Lead Scoring Matrix "
        "interface, allowing administrators to weight factors such as budget sufficiency, academic standing, and document completeness to prioritize "
        "high-conversion leads for counsellors. Both consoles were submitted to Dr. Nadeem Anjum (HOD) for algorithmic review."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '11_lead_scoring_ai.png'),
        "Figure 6.2: AI Lead Scoring Matrix & Factor Configuration Interface"
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '15_ai_programme_matcher.png'),
        "Figure 6.3: AI Programme Eligibility Matcher & Recommendation Console"
    )

    add_day_entry(
        "9 September, 2026 — Wednesday",
        "I developed the Data Quality Dashboard (/data-quality) and Super Admin System Health monitors (/super-admin/dashboard). "
        "The Data Quality dashboard scans the CRM database to detect anomalies such as incomplete student profiles, unassigned leads older "
        "than 48 hours, missing mandatory compliance files, and orphan applications. I submitted the system health dashboards to "
        "Dr. Nadeem Anjum (HOD) who commended the real-time database validation metrics."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '12_data_quality_dashboard.png'),
        "Figure 6.4: Data Quality & Super Admin Health Monitoring Dashboard"
    )

    add_day_entry(
        "10 September, 2026 — Thursday",
        "I verified and enhanced the immutable Audit Log Viewer (/audit-log), which records administrative actions, role impersonation events, "
        "data export requests, and critical record edits to ensure GDPR compliance and institutional auditability. Every action contains an immutable "
        "timestamp, IP address hash, operator identity, and before/after delta values. The audit logging module was submitted to Dr. Nadeem Anjum (HOD) "
        "for regulatory compliance inspection."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'crm_screenshots', '16_audit_log_viewer.png'),
        "Figure 6.5: Immutable System Audit Log Viewer for Regulatory & GDPR Compliance"
    )

    add_day_entry(
        "11 September, 2026 — Friday",
        "To conclude the sixth week, I synthesized the comprehensive end-to-end System Sequence Diagram (SSD), documenting every actor interaction "
        "across the entire EduCRM boundary (Student, Agent, Counsellor, Admissions Officer, Super Admin, Drive API, and Gemini AI API). "
        "I reviewed the complete architectural flow with Mr. Waqar Akram and submitted the diagram and weekly dossier to Dr. Nadeem Anjum (HOD). "
        "Dr. Nadeem Anjum approved the system architecture and authorized final week test automation and handover preparations."
    )

    add_image_box(
        os.path.join(os.getcwd(), 'diagrams_png', 'SSD_System_Sequence_Diagram.png'),
        "Figure 6.6: Comprehensive System Sequence Diagram (SSD) — End-to-End EduCRM Architecture"
    )

    # ----------------------------------------------------
    # WEEK 7: Automated Testing, Optimization & Handover
    # ----------------------------------------------------
    doc.add_page_break()
    add_week_header("Week 7: Automated Testing, Optimization & Final Handover")

    add_day_entry(
        "14 September, 2026 — Monday",
        "At the start of the final week, I focused on comprehensive automated end-to-end testing using the Playwright testing framework. "
        "Working under the direction of Mr. Waqar Akram, I executed automated test suites covering authentication flows, role switching, "
        "student onboarding transitions, and admissions state changes. I verified test cases in tests/e2e/auth.spec.ts and "
        "tests/unit/streamlined-onboarding.test.ts, ensuring that test assertions achieved 100% pass rates across simulated user sessions. "
        "I submitted the automated test execution logs to Dr. Nadeem Anjum (HOD) for verification."
    )

    add_day_entry(
        "15 September, 2026 — Tuesday",
        "Today, I conducted cross-browser compatibility testing and responsive UI verification. I verified that the application renders "
        "flawlessly across Chromium, Mozilla Firefox, and Apple WebKit rendering engines, as well as on various mobile screen viewports. "
        "I eliminated minor CSS layout shifts, polished Tailwind styling utilities, and validated WCAG accessibility contrast ratios across dark "
        "and light theme modes. I submitted the cross-browser audit report to Dr. Nadeem Anjum (HOD) for technical approval."
    )

    add_day_entry(
        "16 September, 2026 — Wednesday",
        "I performed comprehensive performance profiling and bundle optimization. Using Vite's production build analyzer, I reviewed asset "
        "chunking, lazy loading of heavy route components, and code splitting across third-party vendor dependencies. I also verified Firestore "
        "composite indexing rules, ensuring that complex queries on leads, applications, and documents execute with sub-second response times. "
        "I submitted the performance optimization benchmarks to Dr. Nadeem Anjum (HOD) who noted the significant load-time improvements."
    )

    add_day_entry(
        "17 September, 2026 — Thursday",
        "Today, I prepared the staging deployment environment and performed a final data integrity audit. I executed database validation scripts "
        "(scripts/validate-data-integrity.ts) to verify relationship integrity between students, applications, and assigned counsellors. "
        "I ensured that security rules strictly enforce read/write isolation across organizations and that no orphaned documents exist in the database. "
        "The complete staging verification dossier was submitted to Dr. Nadeem Anjum (HOD) for pre-handover evaluation."
    )

    add_day_entry(
        "18 September, 2026 — Friday",
        "I conducted a pre-handover review meeting with Dr. Nadeem Anjum (HOD) at CUST, presenting the entire system documentation, "
        "architectural diagrams, user manuals, and automated test reports. Dr. Nadeem Anjum reviewed the complete 7-week trajectory, praised the "
        "exceptional technical depth, user interface polish, and robust multi-tenant design, and provided final recommendations for the company handover."
    )

    add_day_entry(
        "21 September, 2026 — Monday",
        "I delivered a live, comprehensive demonstration of the finished Education CRM platform to Convo Director Mr. Sanaullah Irfan, "
        "Team Lead Mr. Waqar Akram, and CUST academic supervisor Dr. Nadeem Anjum (HOD). I walked through the complete lifecycle: multi-role authentication, "
        "student onboarding, university discovery, automated lead routing, admissions verification, external agent referral ledgers, and Gemini AI "
        "CV parsing. The presentation was met with enthusiastic acclaim for its enterprise readiness, visual aesthetics, and architectural integrity."
    )

    add_day_entry(
        "22 September, 2026 — Tuesday",
        "On the final day of my 7-week internship, I organized the code repository, consolidated documentation, archived test artifacts, "
        "and completed formal project handover procedures at Convo (Pvt.) Ltd. Following corporate exit clearance, I reported to Capital University "
        "of Science & Technology (CUST) to present the final internship dossier to Dr. Nadeem Anjum (HOD). Dr. Nadeem Anjum conducted the final evaluation, "
        "signed the internship completion certificates, and officially concluded my summer internship."
    )

    # ----------------------------------------------------
    # 6. CONCLUSION
    # ----------------------------------------------------
    doc.add_page_break()
    p_concl_h = doc.add_paragraph()
    p_concl_h.paragraph_format.space_before = Pt(12)
    p_concl_h.paragraph_format.space_after = Pt(16)
    r_ch = p_concl_h.add_run("Conclusion")
    r_ch.font.name = 'Times New Roman'
    r_ch.font.size = Pt(16)
    r_ch.font.bold = True

    p_concl_1 = doc.add_paragraph()
    p_concl_1.paragraph_format.line_spacing = 1.2
    p_concl_1.paragraph_format.space_after = Pt(12)
    p_concl_1.add_run(
        "My seven-week full-stack software engineering internship (3 August, 2026 to 22 September, 2026) at Convo (Pvt.) Ltd., "
        "conducted under the rigorous academic supervision of Dr. Nadeem Anjum (HOD, Department of Software Engineering, CUST), "
        "has been an immensely transformative, challenging, and professionally rewarding journey. Transitioning from academic coursework "
        "at Capital University of Science & Technology into a production-grade enterprise software engineering environment allowed me to "
        "apply theoretical principles of software design, database modeling, and automated testing to a large-scale, mission-critical platform."
    )

    p_concl_2 = doc.add_paragraph()
    p_concl_2.paragraph_format.line_spacing = 1.2
    p_concl_2.paragraph_format.space_after = Pt(12)
    p_concl_2.add_run(
        "The daily guidance and structured milestone evaluations conducted by Dr. Nadeem Anjum (HOD) ensured that every technical deliverable "
        "strictly adhered to software engineering standards, modular architecture, and thorough documentation. In parallel, the industrial mentorship "
        "provided by Director Mr. Sanaullah Irfan and Team Lead Mr. Waqar Akram at Convo expanded my practical capabilities in React 19, "
        "TypeScript, Tailwind CSS v4, and Google Firebase Firestore, while imparting deep insights into multi-tenant SaaS security, asynchronous "
        "state management, and enterprise B2B partner collaboration."
    )

    p_concl_3 = doc.add_paragraph()
    p_concl_3.paragraph_format.line_spacing = 1.2
    p_concl_3.paragraph_format.space_after = Pt(12)
    p_concl_3.add_run(
        "A standout highlight of the internship was integrating cutting-edge artificial intelligence into traditional operational workflows. "
        "Engineering Google Gemini AI integrations for automated resume OCR parsing, intelligent programme eligibility matching, and algorithmic "
        "lead scoring demonstrated firsthand how modern agentic AI can radically enhance business efficiency, reduce manual data entry errors, "
        "and deliver hyper-personalized user experiences."
    )

    p_concl_4 = doc.add_paragraph()
    p_concl_4.paragraph_format.line_spacing = 1.2
    p_concl_4.paragraph_format.space_after = Pt(20)
    p_concl_4.add_run(
        "Beyond technical engineering competencies, this internship honed my professional skills in agile sprint methodologies, git collaboration "
        "workflows, automated browser testing with Playwright, and clear cross-functional reporting. The rigorous dual mentorship of academia "
        "and industry has established a solid, enduring foundation for my future career as a professional software engineer."
    )

    # ----------------------------------------------------
    # 7. REFERENCES
    # ----------------------------------------------------
    doc.add_page_break()
    p_ref_h = doc.add_paragraph()
    p_ref_h.paragraph_format.space_before = Pt(12)
    p_ref_h.paragraph_format.space_after = Pt(16)
    r_rh = p_ref_h.add_run("References")
    r_rh.font.name = 'Times New Roman'
    r_rh.font.size = Pt(16)
    r_rh.font.bold = True

    references = [
        ("1. Capital University of Science and Technology (CUST). ", "SE4100: Internship Program Rules, Regulations, and Documentation Guidelines. Department of Software Engineering, Islamabad, Pakistan."),
        ("2. Convo (Pvt.) Ltd. ", "Internal Engineering Standards, Multi-Tenant SaaS Architecture, and AI Integration Guidelines. Head Office, Islamabad, Pakistan."),
        ("3. EduCRM System Architecture & Flow Specifications. ", "Requirements Traceability Matrix, UML Sequence Diagrams, and Multi-Role Specifications. Convo (Pvt.) Ltd."),
        ("4. IEEE Std 1016-2009. ", "IEEE Standard for Information Technology—Systems Design—Software Design Descriptions. IEEE Computer Society."),
        ("5. ISO/IEC/IEEE 29148:2018. ", "Systems and software engineering — Life cycle processes — Requirements engineering. International Organization for Standardization."),
        ("6. Bass, L., Clements, P., & Kazman, R. ", "Software Architecture in Practice (4th ed.). Addison-Wesley Professional. Multi-tenant Cloud Application Architecture and Role-Based Access Control Patterns."),
        ("7. Google Firebase / Cloud Firestore. ", "NoSQL Cloud Database Design, Real-time Subscriptions, Security Rules, and Multi-Tenant Indexing. Google Cloud & Firebase Documentation (firebase.google.com)."),
        ("8. Google Gemini AI API. ", "Multimodal Large Language Models, Structured Document OCR Parsing, and Contextual Intelligence Integration. Google AI for Developers Documentation (ai.google.dev)."),
        ("9. Playwright Framework. ", "End-to-End Browser Automation, Cross-Browser Testing, and Resilient UI Integration Verification. Microsoft Playwright Documentation (playwright.dev)."),
    ]

    for num_lead, text in references:
        p_r = doc.add_paragraph()
        p_r.paragraph_format.left_indent = Inches(0.3)
        p_r.paragraph_format.first_line_indent = Inches(-0.3)
        p_r.paragraph_format.line_spacing = 1.15
        p_r.paragraph_format.space_after = Pt(8)
        
        r_lead = p_r.add_run(num_lead)
        r_lead.font.bold = True
        r_lead.font.size = Pt(11)
        
        r_body = p_r.add_run(text)
        r_body.font.size = Pt(11)

    # Save to workspace and desktop
    out_workspace = os.path.join(os.getcwd(), 'SAAD_ALI_QURESHI_BSE223018_INTERNSHIP_REPORT.docx')
    doc.save(out_workspace)
    print(f"Report saved to workspace: {out_workspace}")

    desktop_path = r"c:\Users\m\OneDrive\Desktop"
    if os.path.exists(desktop_path):
        out_desktop = os.path.join(desktop_path, 'SAAD_ALI_QURESHI_BSE223018_INTERNSHIP_REPORT.docx')
        shutil.copyfile(out_workspace, out_desktop)
        print(f"Report copied to Desktop: {out_desktop}")

if __name__ == '__main__':
    build_report()

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def add_page_break(doc_or_p):
    """Adds a page break paragraph or break element."""
    p = doc_or_p.add_paragraph() if hasattr(doc_or_p, 'add_paragraph') else doc_or_p
    r = p.add_run()
    r.add_break(docx.enum.text.WD_BREAK.PAGE)
    return p

def set_run_font(run, font_name="Times New Roman", font_size_pt=11, bold=False, italic=False, color_rgb=(0,0,0)):
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)

def set_p_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before_pt=0, space_after_pt=6, line_spacing=1.2):
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.line_spacing = line_spacing

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    set_p_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before_pt=12, space_after_pt=8, line_spacing=1.15)
    r = p.add_run(text)
    set_run_font(r, font_name="Times New Roman", font_size_pt=15, bold=True, color_rgb=(20, 45, 80))
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    set_p_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before_pt=10, space_after_pt=6, line_spacing=1.15)
    r = p.add_run(text)
    set_run_font(r, font_name="Times New Roman", font_size_pt=13, bold=True, color_rgb=(30, 60, 100))
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    set_p_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before_pt=8, space_after_pt=4, line_spacing=1.15)
    r = p.add_run(text)
    set_run_font(r, font_name="Times New Roman", font_size_pt=11.5, bold=True, color_rgb=(40, 70, 110))
    return p

def add_body_paragraph(doc, text, bold_prefix=None, space_after_pt=6):
    p = doc.add_paragraph()
    set_p_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before_pt=0, space_after_pt=space_after_pt, line_spacing=1.2)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, font_name="Times New Roman", font_size_pt=11, bold=True)
    r_body = p.add_run(text)
    set_run_font(r_body, font_name="Times New Roman", font_size_pt=11, bold=False)
    return p

def add_bullet_item(doc, bold_prefix, text, space_after_pt=4):
    p = doc.add_paragraph()
    set_p_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before_pt=0, space_after_pt=space_after_pt, line_spacing=1.18)
    p.paragraph_format.left_indent = Inches(0.25)
    r_bullet = p.add_run("• ")
    set_run_font(r_bullet, font_name="Times New Roman", font_size_pt=11, bold=True, color_rgb=(20, 80, 140))
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, font_name="Times New Roman", font_size_pt=11, bold=True)
    r_body = p.add_run(text)
    set_run_font(r_body, font_name="Times New Roman", font_size_pt=11, bold=False)
    return p

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    set_p_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before_pt=4, space_after_pt=6, line_spacing=1.0)
    p.paragraph_format.left_indent = Inches(0.3)
    pPr = p._p.get_or_add_pPr()
    # Light gray background
    shd_xml = parse_xml('<w:shd ' + nsdecls('w') + ' w:fill="F4F6F8"/>')
    pPr.append(shd_xml)
    pBdr_xml = parse_xml(
        '<w:pBdr ' + nsdecls('w') + '>'
        '<w:left w:val="single" w:sz="18" w:space="8" w:color="10B981"/>'
        '</w:pBdr>'
    )
    pPr.append(pBdr_xml)
    r = p.add_run(code_text)
    set_run_font(r, font_name="Consolas", font_size_pt=9.5, bold=False, color_rgb=(30, 41, 59))
    return p

def add_image_figure(doc, image_path, caption_text, width_inches=6.2):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = parse_xml(
        '<w:tcBorders ' + nsdecls('w') + '>'
        '<w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        '<w:left w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        '<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        '<w:right w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders_xml)
    
    # Internal cell margins
    tcMar_xml = parse_xml(
        '<w:tcMar ' + nsdecls('w') + '>'
        '<w:top w:w="60" w:type="dxa"/>'
        '<w:bottom w:w="60" w:type="dxa"/>'
        '<w:left w:w="60" w:type="dxa"/>'
        '<w:right w:w="60" w:type="dxa"/>'
        '</w:tcMar>'
    )
    tcPr.append(tcMar_xml)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run()
    r.add_picture(image_path, width=Inches(width_inches))
    
    # Caption below table
    p_cap = doc.add_paragraph()
    set_p_format(p_cap, align=WD_ALIGN_PARAGRAPH.CENTER, space_before_pt=4, space_after_pt=10, line_spacing=1.15)
    r_cap = p_cap.add_run(caption_text)
    set_run_font(r_cap, font_name="Times New Roman", font_size_pt=10, bold=False, italic=True, color_rgb=(60, 60, 60))
    return table, p_cap

print("Helper functions defined successfully.")

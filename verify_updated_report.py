import docx

doc = docx.Document(r'c:\Users\m\OneDrive\Desktop\CRM\SAAD_ALI_QURESHI_BSE223018_INTERNSHIP_REPORT.docx')

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

print("\n--- Key Section Headings Found ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt in [
        "Executive Summary", "Acknowledgement", "Introduction", 
        "DAILY ASSIGNMENT TASKS", "Conclusion", "References",
        "Appendix A: Local Setup, Live Deployment & Source Repository",
        "Appendix B: System Dashboards & Core Feature Showcase"
    ] or txt.startswith("Week ") or txt.startswith("B."):
        print(f"P[{i:03d}]: {txt}")

print("\n--- Checking Team Member Mujtaba Zaheer & Registration No. BSE231010 ---")
found_mujtaba = False
for p in doc.paragraphs:
    if "Mujtaba Zaheer" in p.text and "BSE231010" in p.text:
        print(f"Found in paragraph: {p.text[:120]}...")
        found_mujtaba = True
assert found_mujtaba, "Mujtaba Zaheer with BSE231010 not found!"

print("\n--- Checking Executive Summary ---")
found_es = False
for p in doc.paragraphs:
    if "This report documents the software engineering summer internship completed by Saad Ali Qureshi" in p.text:
        print(f"Found Executive Summary text: {p.text[:120]}...")
        found_es = True
assert found_es, "Executive summary text not found!"

print("\n--- Checking Live Deploy Link & Repo ---")
found_live = False
found_repo = False
for p in doc.paragraphs:
    if "https://education-crm-9fee2.web.app" in p.text:
        found_live = True
    if "https://github.com/MujtabaZaheer/CRM.git" in p.text:
        found_repo = True
assert found_live, "Live application URL not found!"
assert found_repo, "GitHub repository URL not found!"
print("Live link and GitHub repository verified!")

print("\n--- Checking Convo mentions outside day-wise assignments ---")
convo_outside = []
is_daily = False
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt == "DAILY ASSIGNMENT TASKS":
        is_daily = True
    elif txt == "Conclusion":
        is_daily = False
    
    if not is_daily and "convo" in txt.lower():
        convo_outside.append((i, txt))

if convo_outside:
    print(f"WARNING: Found {len(convo_outside)} Convo mentions outside daily tasks:")
    for idx, txt in convo_outside:
        print(f"  P[{idx}]: {txt[:80]}")
else:
    print("SUCCESS: Zero mentions of Convo outside day-wise assignments!")

print("\n--- Checking Tables (Existing figures + 5 New Figures in Appendix B) ---")
print(f"Total tables: {len(doc.tables)} (Original had 25, now should have 30)")
assert len(doc.tables) == 30, f"Expected 30 tables, got {len(doc.tables)}"
print("All 30 image tables verified!")

print("\nALL VERIFICATIONS PASSED SUCCESSFULLY!")
